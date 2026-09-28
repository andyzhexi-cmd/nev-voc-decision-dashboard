"""services.report — 导出报告生成服务（纯函数，零 Streamlit 依赖）

设计
----
* 入参是普通数据结构（dict / pandas.DataFrame / VIKORResult），出参是落盘路径（Path）。
* `context_from()` 把「场景 + VIKOR 结果 + 审计 + features 聚合」组装成统一 context，
  三个 build_* 共享同一份 context，视图侧只需一次组装。
* 三种产物：
    - build_excel   多 Sheet Excel（openpyxl，标题行 + 生成时间 + 自适应列宽 + 条件格式）
    - build_markdown 结构化中文 Markdown 报告
    - build_pdf     reportlab 中文 PDF（系统字体 → CID STSong-Light 兜底）
* 依赖缺失时抛 `ReportDependencyError`，由调用方降级为 Markdown 下载。

context 结构（见 `context_from` 的 docstring）
---------------------------------------------
全部为可 pickle 的纯数据，便于 `st.cache_data` 直接作为缓存 key 的一部分。
"""
from __future__ import annotations

import json
import math
import time
import unicodedata
from functools import lru_cache
from pathlib import Path
from typing import Any, Sequence

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
REPORTS_DIR = ROOT / "outputs" / "reports"

KIND_EXT = {"excel": "xlsx", "markdown": "md", "pdf": "pdf"}
KIND_MIME = {
    "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "markdown": "text/markdown; charset=utf-8",
    "pdf": "application/pdf",
}

STATUS_LABEL = {"pass": "✓ 通过", "warn": "△ 待确认", "fail": "✕ 偏差", "info": "ℹ 说明"}
SEVERITY_LABEL = {"high": "高", "medium": "中", "low": "低"}
ISA_LABEL = {"paper": "论文表5.20", "figure": "图5.10", "sentiment": "实测情感"}
QUAD_ORDER = ["改进区", "保持区", "机会区", "低优先级区"]
DEFAULT_TITLE = "智评车行 · 新能源汽车产品改进多属性决策洞察报告"

__all__ = [
    "REPORTS_DIR", "KIND_MIME", "ReportDependencyError",
    "context_from", "sentiment_isa", "isa_conflicts",
    "build_excel", "build_markdown", "build_pdf", "build_payload",
    "new_output_path", "pdf_probe",
]


class ReportDependencyError(RuntimeError):
    """PDF 生成依赖（reportlab / 中文字体）不可用 —— 调用方应降级为 Markdown。"""


# =============================================================== 小工具
def _now() -> str:
    return time.strftime("%Y-%m-%d %H:%M:%S")


def _r(v: Any, n: int = 4) -> Any:
    """四舍五入；非数值原样返回。"""
    try:
        if v is None or (isinstance(v, float) and math.isnan(v)):
            return None
        return round(float(v), n)
    except (TypeError, ValueError):
        return v


def _f(v: Any, n: int = 4, dash: str = "—") -> str:
    if v is None:
        return dash
    if isinstance(v, bool):
        return "是" if v else "否"
    if isinstance(v, (int, np.integer)):
        return f"{int(v):,}"
    if isinstance(v, (float, np.floating)):
        f = float(v)
        if math.isnan(f):
            return dash
        return f"{f:.{n}f}"
    return str(v)


def _disp_width(s: str) -> int:
    """显示宽度：中日韩全角字符按 2 计。"""
    n = 0
    for ch in str(s):
        n += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return n


def _empty_df(note: str = "暂无数据") -> pd.DataFrame:
    return pd.DataFrame({"提示": [note]})


def _json(v: Any) -> str:
    try:
        return json.dumps(v, ensure_ascii=False, default=str)
    except (TypeError, ValueError):
        return str(v)


def _finding_dict(f: Any) -> dict:
    if isinstance(f, dict):
        return dict(f)
    if hasattr(f, "model_dump"):
        return f.model_dump()
    return {
        "id": getattr(f, "id", ""), "severity": getattr(f, "severity", "low"),
        "title": getattr(f, "title", ""), "detail": getattr(f, "detail", ""),
        "status": getattr(f, "status", "info"),
        "numbers": dict(getattr(f, "numbers", {}) or {}),
    }


def _audit_summary(findings: Sequence[dict]) -> dict:
    return {
        "total": len(findings),
        "pass": sum(1 for f in findings if f.get("status") == "pass"),
        "warn": sum(1 for f in findings if f.get("status") == "warn"),
        "fail": sum(1 for f in findings if f.get("status") == "fail"),
        "high": sum(1 for f in findings if f.get("severity") == "high"),
    }


# =============================================================== ISA 辅助
def sentiment_isa(overview: "pd.DataFrame | None") -> dict | None:
    """实测情感口径 ISA：attr_overview 情感均值 → sentiment_to_likert(1-5)。

    overview 为空时返回 None（调用方展示空状态）。
    """
    from core.algorithm.isa import isa_quadrants, sentiment_to_likert
    from core.algorithm.models import baselines

    if overview is None or len(overview) == 0 or "情感均值" not in overview.columns:
        return None
    vals: dict[str, float] = {}
    for _, row in overview.iterrows():
        a = row.get("属性")
        if a is None or (isinstance(a, float) and math.isnan(a)):
            continue
        try:
            vals[str(a)] = float(row["情感均值"])
        except (TypeError, ValueError):
            continue
    if not vals:
        return None
    likert = sentiment_to_likert(vals)
    cfg = baselines()["isa"]
    return isa_quadrants(source="sentiment", satisfaction=likert,
                         importance_mean=float(cfg["importance_mean"]))


def isa_conflicts() -> list[dict]:
    """论文表5.20 与 图5.10 两套满意度口径的冲突明细。"""
    from core.algorithm.isa import isa_quadrants
    from core.algorithm.models import baselines

    b = baselines()
    cfg = b["isa"]
    attrs = list(b["meta"]["attributes"])
    p, f510 = dict(cfg["satisfaction"]), dict(cfg["satisfaction_figure510"])
    qp = {r["属性"]: r["象限"] for r in isa_quadrants(source="paper")["records"]}
    qf = {r["属性"]: r["象限"] for r in isa_quadrants(source="figure")["records"]}
    rows = []
    for a in attrs:
        if abs(float(p[a]) - float(f510[a])) <= 1e-9:
            continue
        rows.append({
            "属性": a,
            "表5.20满意度": float(p[a]),
            "图5.10满意度": float(f510[a]),
            "差值": round(float(f510[a]) - float(p[a]), 3),
            "表5.20象限": qp.get(a, "—"),
            "图5.10象限": qf.get(a, "—"),
            "象限变化": "是" if qp.get(a) != qf.get(a) else "否",
        })
    return rows


# =============================================================== context
def context_from(
    scenario: Any = None,
    result: Any = None,
    findings: Sequence[Any] | None = None,
    overview: "pd.DataFrame | None" = None,
    insight: dict | None = None,
    kpi: dict | None = None,
    filters_desc: str = "",
    data_version: int | None = None,
    isa: dict | None = None,
    isa_source: str = "paper",
    sweep_lambda_res: dict | None = None,
    sweep_v_res: dict | None = None,
    methods: dict | None = None,
    title: str | None = None,
    generated_at: str | None = None,
) -> dict:
    """把 store 侧数据组装成 build_excel / build_markdown / build_pdf 的统一 context。

    参数
    ----
    scenario       core.algorithm.Scenario（缺省用默认场景）
    result         VIKORResult（缺省现算 run_vikor(scenario)）
    findings       list[AuditFinding] 或 list[dict]（缺省现算 run_audit）
    overview       services.features.attr_overview 的 DataFrame（缺省尝试自行加载，失败置空）
    insight / kpi  services.features.insight_cards / kpi_summary 的返回值（缺省置空）
    isa / isa_source  当前 ISA 口径（paper / figure / sentiment）；缺省用论文表5.20
    sweep_lambda_res / sweep_v_res / methods  已算好的扫描与方法对比（缺省现算，避免视图重复计算）
    generated_at   生成时间；缺省不写入 context，由 build_* 在生成时现取（保证缓存 key 稳定）

    返回的 dict 全部为可 pickle 的纯数据（含一个 attr_overview DataFrame）。
    """
    from core.algorithm.audit import run_audit
    from core.algorithm.benchmarks import compare_methods
    from core.algorithm.isa import isa_quadrants
    from core.algorithm.models import Scenario, baselines
    from core.algorithm.sensitivity import (
        compare_to_paper_sensitivity,
        paper_sensitivity_table,
        sweep_lambda,
        sweep_v,
    )
    from core.algorithm.vikor import run_vikor

    b = baselines()
    sc = scenario if scenario is not None else Scenario()
    res = result if result is not None else run_vikor(sc)

    # ---- 属性情感明细（缺省时尽力加载，失败置空，不阻塞导出）
    if overview is None:
        try:
            from services.features import attr_overview
            overview = attr_overview()
        except Exception:                                    # noqa: BLE001
            overview = None
    overview = overview if isinstance(overview, pd.DataFrame) else pd.DataFrame()

    # ---- 审计
    if findings is None:
        try:
            findings = run_audit(sc)
        except Exception:                                    # noqa: BLE001
            findings = []
    f_dicts = [_finding_dict(f) for f in (findings or [])]

    # ---- 论文基准（表5.16/5.17）
    vk = b["vikor"]
    attrs = list(b["meta"]["attributes"])
    paper = {
        "attributes": attrs,
        "S": [float(x) for x in vk["S"]], "R": [float(x) for x in vk["R"]],
        "Q": [float(x) for x in vk["Q"]], "Q_prime": [float(x) for x in vk["Q_prime"]],
        "Pxi": [float(x) for x in vk["Pxi"]],
        "ranking": list(vk["ranking"]),
    }

    # ---- 决策结果（复算 vs 论文，含 Δ）
    cur_rank = {a: i + 1 for i, a in enumerate(res.ranking)}
    pap_rank = {a: i + 1 for i, a in enumerate(paper["ranking"])}
    pmap = {a: k for k, a in enumerate(attrs)}
    rows = []
    for i, a in enumerate(attrs):
        k = pmap[a]
        rows.append({
            "属性": a,
            "复算排名": cur_rank.get(a),
            "论文排名": pap_rank.get(a),
            "排序一致": "是" if cur_rank.get(a) == pap_rank.get(a) else "否",
            "P(x)_复算": _r(res.Pxi[i]),
            "P(x)_论文": paper["Pxi"][k],
            "ΔP(x)": _r(float(res.Pxi[i]) - paper["Pxi"][k]),
            "S_i": _r(res.S[i]), "R_i": _r(res.R[i]), "Q_i": _r(res.Q[i]),
            "Q_论文": paper["Q"][k],
            "ΔQ": _r(float(res.Q[i]) - paper["Q"][k]),
            "Q'_i": _r(res.Q_prime[i]), "σ_i": _r(res.sigma[i]),
        })

    # ---- 权重与中间量
    bnames = b["meta"]["first_level"]
    paper_w = [float(x) for x in b["combined_weights"]["w"]]
    wA = res.diagnostics.get("wA") or [float(x) for x in b["combined_weights"]["wA"]]
    wD = res.diagnostics.get("wD") or [float(x) for x in b["dematel"]["first_weights"]]
    grp = b["meta"]["group_of"]
    sub_paper = b["dematel"]["sub_weights"]
    crit = list(b["meta"]["second_level"])
    crit_names = b["meta"]["second_level_names"]
    reg = b["ideal_solutions"]["regularized"]
    weights = {
        "first": [{
            "一级": bnames[i], "当前综合权重": _r(res.w3[i]), "论文权重": paper_w[i],
            "Δ": _r(float(res.w3[i]) - paper_w[i]),
            "AHP_wA": _r(wA[i]), "DEMATEL_wD": _r(wD[i]),
            "说明": f"式(4.7) w=λ·wD+(1−λ)·wA，当前 λ={sc.lam:g}",
        } for i in range(3)],
        "second": [{
            "指标": c, "名称": crit_names.get(c, c), "一级": grp.get(c, "—"),
            "当前权重": _r(res.w12[i]),
            "DEMATEL二级基准": _r(sub_paper.get(c)),
            "说明": f"按 {sc.weight_mode} 从一级权重展开",
        } for i, c in enumerate(crit)],
        "ideals": {
            "f_star": [float(x) for x in res.f_star],
            "f_minus": [float(x) for x in res.f_minus],
            "paper_f_star": [float(x) for x in reg["f_star"]],
            "paper_f_minus": [float(x) for x in reg["f_minus"]],
            "scope": res.diagnostics.get("ideal_scope", "criterion"),
        },
        "Q_prime": [float(x) for x in res.Q_prime],
        "sigma": [float(x) for x in res.sigma],
        "diagnostics": {k: (v if not isinstance(v, np.ndarray) else v.tolist())
                        for k, v in res.diagnostics.items()},
    }

    # ---- 敏感性
    sw_l = sweep_lambda_res if sweep_lambda_res is not None else sweep_lambda(sc)
    sw_v = sweep_v_res if sweep_v_res is not None else sweep_v(sc)
    paper_sens = paper_sensitivity_table()
    try:
        sens_compare = compare_to_paper_sensitivity(sw_l)
    except Exception:                                        # noqa: BLE001
        sens_compare = []

    # ---- 方法对比
    cm = methods if methods is not None else compare_methods(sc)
    m_rows = [dict(r) for r in cm.get("rows", [])]
    m_scores = {}
    for m in cm.get("methods", []):
        arr = np.asarray(m.get("scores", []), dtype=float)
        m_scores[m["name"]] = [round(float(x), 4) for x in arr]

    # ---- ISA（三口径）
    variants: dict[str, dict] = {
        "paper": isa_quadrants(source="paper"),
        "figure": isa_quadrants(source="figure"),
    }
    s_isa = sentiment_isa(overview)
    if s_isa is not None:
        variants["sentiment"] = s_isa
    current = isa if isa is not None else variants.get(isa_source)
    if current is None:
        current = variants["paper"]
        isa_source = "paper"

    ctx = {
        "title": title or DEFAULT_TITLE,
        "paper_title": b["meta"].get("paper_title", ""),
        "thesis": b["meta"].get("thesis", ""),
        "data_version": int(data_version) if data_version is not None else None,
        "filters_desc": filters_desc or "全量数据（未筛选）",
        "scenario": {
            "mode": sc.mode, "lam": float(sc.lam), "v": float(sc.v),
            "ideal_strategy": sc.ideal_strategy, "weight_mode": sc.weight_mode,
            "prob_completion": sc.prob_completion, "direction": sc.direction,
            "use_paper_weights": bool(sc.use_paper_weights),
            "mode_label": "论文校准模式" if sc.mode == "calibrated" else "在线复算模式",
        },
        "kpi": dict(kpi or {}),
        "insight": dict(insight or {}),
        "results": {
            "rows": rows,
            "attributes": attrs,
            "ranking": list(res.ranking),
            "Q_order": list(res.Q_order),
            "paper_ranking": list(paper["ranking"]),
            "prob_sum_deviation": float(getattr(res, "prob_sum_deviation", 0.0)),
            "diagnostics": weights["diagnostics"],
        },
        "paper": paper,
        "weights": weights,
        "sensitivity": {
            "lambda": {k: sw_l[k] for k in ("lambdas", "Pxi", "ranks", "orders",
                                            "stability", "base_order", "stable")},
            "v": {k: sw_v[k] for k in ("vs", "Pxi", "orders", "stability",
                                       "base_order", "stable")},
            "paper": paper_sens,
            "compare": sens_compare,
        },
        "methods": {
            "rows": m_rows, "scores": m_scores,
            "top1_agreement": cm.get("top1_agreement", "—"),
            "paper": cm.get("paper", {}),
        },
        "isa": current,
        "isa_source": isa_source,
        "isa_source_label": ISA_LABEL.get(isa_source, isa_source),
        "isa_variants": variants,
        "isa_conflict": isa_conflicts(),
        "audit": {"summary": _audit_summary(f_dicts), "findings": f_dicts},
        "attr_detail": overview,
    }
    if generated_at:
        ctx["generated_at"] = generated_at
    return ctx


# =============================================================== DataFrame 工厂
def _ts(ctx: dict) -> str:
    return ctx.get("generated_at") or _now()


def _df_summary(ctx: dict) -> pd.DataFrame:
    sc = ctx.get("scenario", {}) or {}
    res = ctx.get("results", {}) or {}
    kpi = ctx.get("kpi", {}) or {}
    insight = ctx.get("insight", {}) or {}
    audit = ctx.get("audit", {}) or {}
    summ = audit.get("summary", {}) or {}
    rows_r = res.get("rows", []) or []
    ranking = res.get("ranking", []) or []
    paper_rank = res.get("paper_ranking", []) or []
    sens = (ctx.get("sensitivity", {}) or {}).get("lambda", {}) or {}
    taus = sens.get("stability") or [1.0]

    items: list[tuple[str, str, str]] = [
        ("报告标题", ctx.get("title", DEFAULT_TITLE), "洞察与报告 · 一键导出"),
        ("生成时间", _ts(ctx), "服务器本地时间"),
        ("论文基准", f"{ctx.get('paper_title', '')}", ctx.get("thesis", "")),
        ("数据口径", ctx.get("filters_desc", "全量数据（未筛选）"),
         f"数据版本 {ctx.get('data_version') if ctx.get('data_version') is not None else '—'}"),
        ("场景参数",
         f"{sc.get('mode_label', '—')} · λ={sc.get('lam', '—'):g} · v={sc.get('v', '—'):g}"
         if isinstance(sc.get("lam"), (int, float)) else "—",
         f"理想解={sc.get('ideal_strategy', '—')} · 权重={sc.get('weight_mode', '—')} · "
         f"效用方向={sc.get('direction', '—')} · 概率补全={sc.get('prob_completion', '—')}"),
    ]
    if kpi:
        items += [
            ("有效评论数", f"{int(kpi.get('n_kept', 0)):,}", f"保留率 {kpi.get('keep_rate', '—')}%"),
            ("筛选后情感均值", _f(kpi.get("mean_sentiment")),
             f"论文基准 {_f(kpi.get('paper_mean_sentiment'))}"),
            ("正/负评论", f"{int(kpi.get('pos', 0)):,} / {int(kpi.get('neg', 0)):,}",
             f"正负比 {kpi.get('pos_ratio', '—')}%"),
        ]
    if insight.get("headline"):
        items.append(("数据头条", str(insight["headline"]), "由 features.insight_cards 自动生成"))
    if ranking and rows_r:
        top = rows_r[[r["属性"] for r in rows_r].index(ranking[0])] if ranking[0] in [r["属性"] for r in rows_r] else rows_r[0]
        items += [
            ("复算 Top1 属性", f"{ranking[0]}（P(x)={_f(top.get('P(x)_复算'))}）",
             f"完整排序：{' > '.join(ranking)}"),
            ("论文 Top1 属性", f"{paper_rank[0] if paper_rank else '—'}",
             f"论文排序：{' > '.join(paper_rank) if paper_rank else '—'}"),
            ("排序一致性",
             "完全一致" if ranking == paper_rank else "部分不一致（详见决策结果表）",
             f"Kendall τ={_kendall(ranking, paper_rank):.4f}"),
        ]
    if sens:
        items.append(("λ 扫描稳定性",
                      "稳定（7 个 λ 点排序一致）" if sens.get("stable") else "存在排序翻转",
                      f"最低 Kendall τ={min(taus):.4f}"))
    if summ:
        items.append(("一致性审计",
                      f"共 {summ.get('total', 0)} 项 · 通过 {summ.get('pass', 0)} · "
                      f"待确认 {summ.get('warn', 0)} · 偏差 {summ.get('fail', 0)}",
                      f"高严重度 {summ.get('high', 0)} 项；问题来自论文数据本身，系统不做静默修正"))
    isa = ctx.get("isa", {}) or {}
    if isa.get("counts"):
        imp = [r["属性"] for r in isa.get("records", []) if r.get("象限") == "改进区"]
        items.append((f"ISA 当前口径（{ctx.get('isa_source_label', '论文表5.20')}）",
                      f"改进区：{'、'.join(imp) if imp else '无'}",
                      f"重要性均值 {_f(isa.get('importance_mean'))} · "
                      f"满意度均值 {_f(isa.get('satisfaction_mean'))}"))
    return pd.DataFrame(items, columns=["项目", "数值", "说明"])


def _kendall(a: Sequence[str], b: Sequence[str]) -> float:
    from core.algorithm.sensitivity import kendall_tau
    try:
        return float(kendall_tau(list(a), list(b)))
    except Exception:                                        # noqa: BLE001
        return 1.0


def _df_results(ctx: dict) -> pd.DataFrame:
    rows = (ctx.get("results", {}) or {}).get("rows", []) or []
    if not rows:
        return _empty_df("暂无决策结果")
    return pd.DataFrame(rows)


def _df_weights(ctx: dict) -> pd.DataFrame:
    w = ctx.get("weights", {}) or {}
    sc = ctx.get("scenario", {}) or {}
    attrs = (ctx.get("results", {}) or {}).get("attributes", []) or []
    out: list[dict] = []
    for r in w.get("first", []) or []:
        out.append({"分组": "一级综合权重", "代码/项目": r.get("一级"), "名称": "期望值/敏感值/吸引值",
                    "当前值": r.get("当前综合权重"), "论文基准": r.get("论文权重"),
                    "Δ": r.get("Δ"), "说明": r.get("说明")})
        out.append({"分组": "一级分解 · AHP", "代码/项目": r.get("一级"), "名称": "表5.8 主观权重 wA",
                    "当前值": r.get("AHP_wA"), "论文基准": None, "Δ": None, "说明": "λ=0 端点权重"})
        out.append({"分组": "一级分解 · DEMATEL", "代码/项目": r.get("一级"), "名称": "表5.11 关联权重 wD",
                    "当前值": r.get("DEMATEL_wD"), "论文基准": None, "Δ": None, "说明": "λ=1 端点权重"})
    for r in w.get("second", []) or []:
        out.append({"分组": "二级指标权重", "代码/项目": r.get("指标"), "名称": r.get("名称"),
                    "当前值": r.get("当前权重"), "论文基准": r.get("DEMATEL二级基准"),
                    "Δ": None, "说明": f"一级 {r.get('一级')} · {r.get('说明')}"})
    ideals = w.get("ideals", {}) or {}
    f_star = ideals.get("f_star", []) or []
    f_minus = ideals.get("f_minus", []) or []
    p_star = ideals.get("paper_f_star", []) or []
    p_minus = ideals.get("paper_f_minus", []) or []
    scope = ideals.get("scope", "criterion")
    for i, a in enumerate(attrs):
        out.append({"分组": "理想解 f*", "代码/项目": a, "名称": "正理想解",
                    "当前值": _r(f_star[i]) if i < len(f_star) else None,
                    "论文基准": _r(p_star[i]) if i < len(p_star) else None,
                    "Δ": None,
                    "说明": "按规范化矩阵列极值（指标列口径）" if scope == "criterion"
                            else "情感理想解（表5.15 口径）"})
        out.append({"分组": "理想解 f-", "代码/项目": a, "名称": "负理想解",
                    "当前值": _r(f_minus[i]) if i < len(f_minus) else None,
                    "论文基准": _r(p_minus[i]) if i < len(p_minus) else None,
                    "Δ": None,
                    "说明": "按规范化矩阵列极值（指标列口径）" if scope == "criterion"
                            else "情感理想解（表5.15 口径）"})
    qp = w.get("Q_prime", []) or []
    sg = w.get("sigma", []) or []
    for i, a in enumerate(attrs):
        out.append({"分组": "区间参数 Q'", "代码/项目": a, "名称": "折衷评价值正态区间中点",
                    "当前值": _r(qp[i]) if i < len(qp) else None, "论文基准": None, "Δ": None,
                    "说明": "式(4.11) Q 的正态区间表示"})
        out.append({"分组": "区间参数 σ", "代码/项目": a, "名称": "区间标准差",
                    "当前值": _r(sg[i]) if i < len(sg) else None, "论文基准": None, "Δ": None,
                    "说明": "σ=max((Q_hi−Q_lo)/3.92, 0.06)"})
    diag = w.get("diagnostics", {}) or {}
    out.append({"分组": "诊断", "代码/项目": "ideal_scope", "名称": "理想解口径",
                "当前值": str(diag.get("ideal_scope", "—")), "论文基准": None, "Δ": None,
                "说明": "criterion=按规范化矩阵列极值 / attribute=按属性情感理想解"})
    out.append({"分组": "诊断", "代码/项目": "weight_source", "名称": "权重来源",
                "当前值": str(diag.get("weight_source", "—")), "论文基准": None, "Δ": None,
                "说明": f"使用论文权重={sc.get('use_paper_weights', '—')}"})
    out.append({"分组": "诊断", "代码/项目": "ahp_CR", "名称": "AHP 一致性比例",
                "当前值": _r(diag.get("ahp_CR")), "论文基准": 0.0846, "Δ": None,
                "说明": "CR<0.1 视为通过" if diag.get("ahp_consistent") else "CR≥0.1，一致性不足（论文矩阵本身问题）"})
    out.append({"分组": "诊断", "代码/项目": "prob_incomplete", "名称": "Σp<1 的格子数",
                "当前值": int(round(float((ctx.get("results", {}) or {}).get("prob_sum_deviation", 0.0) or 0.0))),
                "论文基准": 0, "Δ": None,
                "说明": f"表5.13 论文原文如此，按 {sc.get('prob_completion', '—')} 策略补全"})
    return pd.DataFrame(out, columns=["分组", "代码/项目", "名称", "当前值", "论文基准", "Δ", "说明"])


def _df_sensitivity(ctx: dict) -> pd.DataFrame:
    sens = ctx.get("sensitivity", {}) or {}
    lam_s = sens.get("lambda", {}) or {}
    v_s = sens.get("v", {}) or {}
    paper = sens.get("paper", {}) or {}
    p_lams = [float(x) for x in (paper.get("lambdas") or [])]
    p_px = paper.get("Pxi", {}) or {}
    out: list[dict] = []

    def _push(param: str, xs, sw, ranks_key: str, with_paper: bool):
        attrs = [a for a in (sw.get("Pxi") or {})]
        for i, x in enumerate(xs):
            tau = (sw.get("stability") or [None] * len(xs))[i]
            for a in attrs:
                val = sw["Pxi"][a][i]
                rank = (sw.get("ranks") or {}).get(a, [None] * len(xs))[i] \
                    if ranks_key == "ranks" and a in (sw.get("ranks") or {}) else None
                pv = None
                if with_paper:
                    for j, pl in enumerate(p_lams):
                        if abs(pl - float(x)) < 1e-9:
                            pv = (p_px.get(a) or [None] * 7)[j]
                            break
                out.append({
                    "参数": param, "取值": float(x), "Kendall_τ": _r(tau),
                    "属性": a, "P(x)_复算": _r(val), "排名": rank,
                    "论文P(x)": _r(pv) if pv is not None else None,
                    "Δ_vs论文": _r(float(val) - float(pv)) if pv is not None else None,
                })

    _push("λ", lam_s.get("lambdas") or [], lam_s, "ranks", True)
    _push("v", v_s.get("vs") or [], v_s, "ranks", False)
    if not out:
        return _empty_df("暂无敏感性扫描结果")
    return pd.DataFrame(out, columns=["参数", "取值", "Kendall_τ", "属性", "P(x)_复算",
                                      "排名", "论文P(x)", "Δ_vs论文"])


def _df_methods(ctx: dict) -> pd.DataFrame:
    m = ctx.get("methods", {}) or {}
    rows = m.get("rows", []) or []
    scores = m.get("scores", {}) or {}
    attrs = (ctx.get("results", {}) or {}).get("attributes", []) or []
    if not rows:
        return _empty_df("暂无方法对比结果")
    out = []
    for r in rows:
        name = r.get("方法", "")
        rec = {
            "方法": name,
            "复算排序": r.get("排序", ""),
            "论文排序": r.get("论文排序", ""),
            "是否一致": "是" if r.get("与论文一致") else "否",
        }
        cur = str(r.get("排序", "")).split(" > ")
        pap = str(r.get("论文排序", "")).split(" > ")
        rec["Top1复算"] = cur[0] if cur else "—"
        rec["Top1论文"] = pap[0] if pap else "—"
        rec["Top1一致"] = "是" if cur and pap and cur[0] == pap[0] else "否"
        for i, a in enumerate(attrs):
            s = scores.get(name) or []
            rec[f"得分·{a}"] = _r(s[i]) if i < len(s) else None
        out.append(rec)
    cols = ["方法", "复算排序", "论文排序", "是否一致", "Top1复算", "Top1论文", "Top1一致"] + \
           [f"得分·{a}" for a in attrs]
    return pd.DataFrame(out, columns=cols)


def _df_isa(ctx: dict) -> pd.DataFrame:
    variants = ctx.get("isa_variants", {}) or {}
    current = ctx.get("isa_source", "paper")
    paper_quad = {r["属性"]: r["象限"]
                  for r in (variants.get("paper") or ctx.get("isa") or {}).get("records", [])}
    out: list[dict] = []
    for src in ("paper", "figure", "sentiment"):
        v = variants.get(src)
        if not v:
            continue
        for r in v.get("records", []):
            out.append({
                "口径": ISA_LABEL.get(src, src) + ("（当前）" if src == current else ""),
                "属性": r.get("属性"), "重要性": r.get("重要性"), "满意度": r.get("满意度"),
                "象限": r.get("象限"),
                "论文口径象限": paper_quad.get(r.get("属性"), "—"),
                "象限一致": "是" if paper_quad.get(r.get("属性")) == r.get("象限") else "否",
                "Δ重要性": r.get("Δ重要性"), "Δ满意度": r.get("Δ满意度"),
            })
    if not out:
        return _empty_df("暂无 ISA 数据")
    return pd.DataFrame(out, columns=["口径", "属性", "重要性", "满意度", "象限",
                                      "论文口径象限", "象限一致", "Δ重要性", "Δ满意度"])


def _df_audit(ctx: dict) -> pd.DataFrame:
    audit = ctx.get("audit", {}) or {}
    fs = audit.get("findings", []) or []
    if not fs:
        return _empty_df("暂无审计发现")
    out = [{
        "编号": f.get("id", ""),
        "严重度": SEVERITY_LABEL.get(str(f.get("severity")), str(f.get("severity"))),
        "状态": STATUS_LABEL.get(str(f.get("status")), str(f.get("status"))),
        "标题": f.get("title", ""),
        "说明": f.get("detail", ""),
        "关键数字": _json(f.get("numbers", {})),
    } for f in fs]
    return pd.DataFrame(out, columns=["编号", "严重度", "状态", "标题", "说明", "关键数字"])


def _df_attr(ctx: dict) -> pd.DataFrame:
    ov = ctx.get("attr_detail")
    if not isinstance(ov, pd.DataFrame) or ov.empty:
        return _empty_df("暂无属性情感明细（请先完成数据流水线或调整筛选）")
    df = ov.copy()
    if "情感均值" in df.columns:
        df["情感李克特(1-5)"] = df["情感均值"].apply(
            lambda x: _r(1.0 + 4.0 * (float(np.clip(x, -1.0, 1.0) + 1.0) / 2.0), 3))
    return df


def _excel_sheets(ctx: dict) -> list[tuple[str, str, pd.DataFrame, dict]]:
    """(sheet 名, 副标题, DataFrame, 样式选项)。"""
    return [
        ("执行摘要", "关键指标与结论一览", _df_summary(ctx), {}),
        ("决策结果", "当前复算 vs 论文表5.16/5.17 基准（Δ 列 = 复算 − 论文）", _df_results(ctx),
         {"scale_cols": ["P(x)_复算", "ΔP(x)", "ΔQ"], "fill_cols": ["排序一致"]}),
        ("权重与中间量", "一级/二级权重、理想解、区间参数与诊断", _df_weights(ctx),
         {"scale_cols": ["当前值"], "fill_cols": []}),
        ("敏感性λ与v", "λ 融合系数与 v 决策偏好扫描（含 Kendall τ 与论文表5.18 对照）",
         _df_sensitivity(ctx), {"scale_cols": ["P(x)_复算"], "fill_cols": []}),
        ("方法对比", "论文表5.19：PLTS-VIKOR / 传统VIKOR / TOPSIS / 前景理论",
         _df_methods(ctx), {"scale_cols": [], "fill_cols": ["是否一致", "Top1一致"]}),
        ("ISA象限", "论文表5.20 / 图5.10 / 实测情感 三口径对照", _df_isa(ctx),
         {"scale_cols": [], "fill_cols": ["象限", "象限一致"]}),
        ("审计发现", "以下问题来自论文数据本身，系统不做静默修正", _df_audit(ctx),
         {"scale_cols": [], "fill_cols": ["状态"]}),
        ("属性情感明细", "筛选后评论的属性情感统计（李克特列为式换算结果）", _df_attr(ctx),
         {"scale_cols": ["情感均值", "Δvs论文"], "fill_cols": []}),
    ]


# =============================================================== Excel
_STATUS_FILL = {
    "✓ 通过": "C6EFCE", "是": "C6EFCE", "通过": "C6EFCE",
    "△ 待确认": "FFEB9C", "否": "FFEB9C", "待确认": "FFEB9C",
    "✕ 偏差": "FFC7CE", "偏差": "FFC7CE",
    "ℹ 说明": "E7E6E6",
}
_QUAD_FILL = {"改进区": "FFC7CE", "保持区": "C6EFCE", "机会区": "DDEBF7", "低优先级区": "E7E6E6"}


def _write_sheet(writer, sheet: str, title: str, subtitle: str, df: pd.DataFrame,
                 ts: str, scale_cols: Sequence[str] = (), fill_cols: Sequence[str] = ()) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    df = df.copy()
    if df is None or df.empty:
        df = _empty_df()
    for c in df.columns:
        if df[c].dtype == object:
            df[c] = df[c].where(pd.notna(df[c]), None)
    df.to_excel(writer, sheet_name=sheet, index=False, startrow=3)
    ws = writer.sheets[sheet]

    # 标题与生成时间
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14, color="1F2937")
    ws.cell(row=2, column=1,
            value=f"生成时间：{ts} · {subtitle}").font = Font(size=9, italic=True, color="6B7280")
    ws.freeze_panes = "A5"

    n_rows = len(df) + 4          # 表头在第 4 行
    n_cols = len(df.columns)
    # 表头样式
    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=4, column=j)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # 列宽 + 数字格式
    for j, col in enumerate(df.columns, start=1):
        width = _disp_width(str(col))
        for i in range(len(df)):
            v = df.iloc[i, j - 1]
            if isinstance(v, (float, np.floating)):
                width = max(width, min(_disp_width(f"{float(v):.4f}") + 2, 60))
            else:
                width = max(width, min(_disp_width("" if v is None else str(v)) + 2, 60))
        ws.column_dimensions[get_column_letter(j)].width = max(9, min(width, 60))
    for row in ws.iter_rows(min_row=5, max_row=n_rows, max_col=n_cols):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
            elif isinstance(cell.value, str) and len(cell.value) > 40:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    # 条件格式
    try:
        from openpyxl.formatting.rule import ColorScaleRule
        for col in scale_cols:
            if col not in df.columns:
                continue
            letter = get_column_letter(list(df.columns).index(col) + 1)
            ws.conditional_formatting.add(
                f"{letter}5:{letter}{n_rows}",
                ColorScaleRule(start_type="min", start_color="F8696B",
                               mid_type="percentile", mid_value=50, mid_color="FFEB84",
                               end_type="max", end_color="63BE7B"))
    except Exception:                                        # noqa: BLE001
        pass
    # 状态 / 象限填色
    for col in fill_cols:
        if col not in df.columns:
            continue
        j = list(df.columns).index(col) + 1
        for i in range(len(df)):
            v = df.iloc[i, j - 1]
            key = str(v)
            color = _QUAD_FILL.get(key) or _STATUS_FILL.get(key)
            if color:
                ws.cell(row=5 + i, column=j).fill = PatternFill("solid", fgColor=color)
    if n_rows >= 4:
        ws.auto_filter.ref = f"A4:{get_column_letter(n_cols)}{n_rows}"


def build_excel(context: dict, out_path: Path) -> Path:
    """多 Sheet Excel 导出（pandas.ExcelWriter + openpyxl）。返回落盘路径。"""
    try:
        from openpyxl import load_workbook  # noqa: F401  (探测依赖)
        writer_engine = "openpyxl"
    except ImportError as e:                                  # pragma: no cover
        raise ReportDependencyError(f"openpyxl 不可用，无法生成 Excel：{e}") from e

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    ts = _ts(context)
    title = context.get("title", DEFAULT_TITLE)

    with pd.ExcelWriter(out_path, engine=writer_engine) as writer:
        for sheet, subtitle, df, opts in _excel_sheets(context):
            _write_sheet(writer, sheet, f"{title} · {sheet}", subtitle, df, ts,
                         scale_cols=opts.get("scale_cols", ()),
                         fill_cols=opts.get("fill_cols", ()))
    return out_path


# =============================================================== Markdown
def _md_table(df: pd.DataFrame) -> str:
    if df is None or len(df) == 0:
        return "_（暂无数据）_\n"
    df = df.head(200)

    def cell(v: Any) -> str:
        if v is None:
            return "—"
        if isinstance(v, bool):
            return "是" if v else "否"
        if isinstance(v, (float, np.floating)):
            f = float(v)
            return "—" if math.isnan(f) else f"{f:.4f}"
        if isinstance(v, (int, np.integer)):
            return str(int(v))
        s = str(v).replace("|", "／").replace("\n", " ")
        return s if len(s) <= 160 else s[:157] + "…"

    cols = [str(c) for c in df.columns]
    head = "| " + " | ".join(cols) + " |"
    sep = "| " + " | ".join("---" for _ in cols) + " |"
    body = ["| " + " | ".join(cell(df.iloc[i, j]) for j in range(len(cols))) + " |"
            for i in range(len(df))]
    return "\n".join([head, sep] + body) + "\n"


def build_markdown(context: dict, out_path: Path) -> Path:
    """结构化中文 Markdown 报告：背景与口径 → 执行摘要 → … → 附录。"""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    ts = _ts(context)
    sc = context.get("scenario", {}) or {}
    res = context.get("results", {}) or {}
    audit = context.get("audit", {}) or {}
    summ = audit.get("summary", {}) or {}
    sens = context.get("sensitivity", {}) or {}
    lam_s = sens.get("lambda", {}) or {}
    v_s = sens.get("v", {}) or {}
    isa = context.get("isa", {}) or {}
    kpi = context.get("kpi", {}) or {}
    insight = context.get("insight", {}) or {}
    rows_r = res.get("rows", []) or []
    ranking = res.get("ranking", []) or []
    paper_rank = res.get("paper_ranking", []) or []

    L: list[str] = []
    A = L.append
    A(f"# {context.get('title', DEFAULT_TITLE)}")
    A("")
    A(f"> 生成时间：{ts}　|　论文基准：{context.get('paper_title', '')}（{context.get('thesis', '')}）")
    A("")

    # 一、背景与口径
    A("## 一、背景与口径")
    A("")
    A("本报告由「智评车行 · 新能源汽车产品改进多属性决策平台」自动生成：以论文《"
      f"{context.get('paper_title', '')}》的表 5.7–5.20 为基准口径，"
      "对评论情感分析结果与 PLTS-VIKOR 多属性决策结果进行复算、对照与审计，"
      "并给出改进优先级建议。**所有已知的论文内部口径冲突与不可复算项均在「八、审计与局限」中显式列出，"
      "系统不做静默修正。**")
    A("")
    A("| 口径项 | 取值 | 说明 |")
    A("| --- | --- | --- |")
    A(f"| 计算模式 | {sc.get('mode_label', '—')} | 校准模式以论文权重与基准为主口径 |")
    A(f"| 融合系数 λ | {sc.get('lam', '—')} | 式(4.7) w = λ·wD + (1−λ)·wA |")
    A(f"| 决策偏好 v | {sc.get('v', '—')} | 式(4.11) 群体效用与个体遗憾的相对权重 |")
    A(f"| 效用方向 | {sc.get('direction', '—')} | 达标度口径（论文 §4.3）：S/Q 越大越重要，P(x) 降序即重要性排序 |")
    A(f"| 理想解口径 | {sc.get('ideal_strategy', '—')} | 按规范化矩阵列极值 / 按属性情感理想解（表5.15） |")
    A(f"| 二级权重展开 | {sc.get('weight_mode', '—')} | 一级权重按组内份额展开到 12 个二级指标（Σ=1） |")
    A(f"| PLTS 概率补全 | {sc.get('prob_completion', '—')} | 表5.13 存在 Σp<1 的格子，按此策略补全 |")
    A(f"| 数据范围 | {context.get('filters_desc', '全量数据（未筛选）')} | "
      f"数据版本 {context.get('data_version') if context.get('data_version') is not None else '—'} |")
    A("")

    # 二、执行摘要
    A("## 二、执行摘要")
    A("")
    if insight.get("headline"):
        A(f"**{insight['headline']}**")
        A("")
    if kpi:
        A(_md_table(pd.DataFrame([
            {"指标": "有效评论数", "数值": f"{int(kpi.get('n_kept', 0)):,}",
             "备注": f"保留率 {kpi.get('keep_rate', '—')}%"},
            {"指标": "筛选后情感均值", "数值": _f(kpi.get("mean_sentiment")),
             "备注": f"论文基准 {_f(kpi.get('paper_mean_sentiment'))}"},
            {"指标": "正/负评论", "数值": f"{int(kpi.get('pos', 0)):,} / {int(kpi.get('neg', 0)):,}",
             "备注": f"正负比 {kpi.get('pos_ratio', '—')}%"},
            {"指标": "覆盖平台/品牌/车型",
             "数值": f"{kpi.get('n_platforms', '—')} / {kpi.get('n_brands', '—')} / {kpi.get('n_models', '—')}",
             "备注": "原始评论 " + f"{int(kpi.get('n_raw', 0)):,} 条"},
        ])))
        A("")
    top3 = "、".join(f"{i + 1}. **{a}**" for i, a in enumerate(ranking[:3])) if ranking else "—"
    A(f"- 复算重要性排序（P(x) 降序）：{top3}")
    A(f"- 论文排序：{' > '.join(paper_rank) if paper_rank else '—'}")
    A(f"- 排序一致性：{'完全一致' if ranking == paper_rank else '**部分不一致**（见决策结果与审计）'}"
      f"（Kendall τ = {_kendall(ranking, paper_rank):.4f}）")
    if summ:
        A(f"- 一致性审计：共 {summ.get('total', 0)} 项，通过 {summ.get('pass', 0)} 项、"
          f"待确认 {summ.get('warn', 0)} 项、偏差 {summ.get('fail', 0)} 项（高严重度 {summ.get('high', 0)} 项）")
    if lam_s.get("stability"):
        A(f"- λ 扫描稳健性：{'排序稳定' if lam_s.get('stable') else '**存在排序翻转**'}，"
          f"最低 Kendall τ = {min(lam_s['stability']):.4f}")
    A("")

    # 三、决策结果与排序
    A("## 三、决策结果与排序")
    A("")
    A("口径：P(x) 由式(4.12)(4.13) 折衷可能度矩阵求和得到，**越大越重要**；Δ 列为当前复算与论文基准的差值。")
    A("")
    A(_md_table(pd.DataFrame(rows_r)))
    A("")
    A(f"- 复算排序：{' > '.join(ranking) if ranking else '—'}")
    A(f"- 论文排序（表5.17）：{' > '.join(paper_rank) if paper_rank else '—'}")
    A(f"- Q 升序（传统 VIKOR 优劣口径）：{' > '.join(res.get('Q_order', []) or []) or '—'}")
    A("")

    # 四、权重体系
    A("## 四、权重体系")
    A("")
    A(_md_table(pd.DataFrame((context.get("weights", {}) or {}).get("first", []) or [])))
    A("")
    A(_md_table(pd.DataFrame((context.get("weights", {}) or {}).get("second", []) or [])))
    A("")
    A(f"- 理想解口径：{(context.get('weights', {}) or {}).get('ideals', {}).get('scope', '—')}；"
      "详细理想解、Q′ 与 σ 见 Excel「权重与中间量」Sheet。")
    A("")

    # 五、敏感性与稳健性
    A("## 五、敏感性与稳健性")
    A("")
    if lam_s.get("lambdas"):
        lam_df = pd.DataFrame(lam_s["Pxi"])
        lam_df.insert(0, "λ", lam_s["lambdas"])
        lam_df.insert(len(lam_df.columns), "Kendall_τ",
                      [_r(t) for t in (lam_s.get("stability") or [None] * len(lam_df))])
        A("**λ 敏感性（复算 P(x)；论文表5.18 对照见 Excel「敏感性λ与v」）**")
        A("")
        A(_md_table(lam_df))
        A("")
    if v_s.get("vs"):
        v_df = pd.DataFrame(v_s["Pxi"])
        v_df.insert(0, "v", v_s["vs"])
        v_df.insert(len(v_df.columns), "Kendall_τ",
                    [_r(t) for t in (v_s.get("stability") or [None] * len(v_df))])
        A("**v 敏感性（复算 P(x)）**")
        A("")
        A(_md_table(v_df))
        A("")
    if sens.get("compare"):
        A("**复算 vs 论文表5.18 对照**")
        A("")
        A(_md_table(pd.DataFrame(sens["compare"])))
        A("")

    # 六、方法对比
    A("## 六、方法对比")
    A("")
    m = context.get("methods", {}) or {}
    A(f"Top1 一致率：**{m.get('top1_agreement', '—')}**（复算 Top1 与论文表5.19 Top1 相同的方法数 / 方法总数）")
    A("")
    A(_md_table(pd.DataFrame(_df_methods(context))))
    A("")

    # 七、ISA 改进优先级
    A("## 七、ISA 改进优先级")
    A("")
    A(f"当前口径：**{context.get('isa_source_label', '论文表5.20')}**；"
      f"重要性均值 {_f(isa.get('importance_mean'))}、满意度均值 {_f(isa.get('satisfaction_mean'))}。")
    A("")
    A(_md_table(pd.DataFrame(isa.get("records", []) or [])))
    A("")
    for q in QUAD_ORDER:
        recs = [r for r in isa.get("records", []) if r.get("象限") == q]
        if not recs:
            continue
        A(f"**{q}（{len(recs)} 项）**")
        for r in recs:
            A(f"- {r.get('属性')}：重要性 {_f(r.get('重要性'))}、满意度 {_f(r.get('满意度'))} —— "
              f"{_advice_body(r)}")
        A("")
    conflict = context.get("isa_conflict", []) or []
    if conflict:
        A("> ⚠ **口径冲突提示**：论文表5.20 与图5.10/正文对同一属性给出不同满意度，"
          "象限归属可能变化（明细见下表与审计项 ISA_SATISFACTION_CONFLICT）。")
        A("")
        A(_md_table(pd.DataFrame(conflict)))
        A("")

    # 八、审计与局限
    A("## 八、审计与局限")
    A("")
    A("> **以下问题来自论文数据本身，系统不做静默修正**：每条核查都附带实时计算出的数字，便于逐项复核。")
    A("")
    for f in audit.get("findings", []) or []:
        st_label = STATUS_LABEL.get(str(f.get("status")), str(f.get("status")))
        sev = SEVERITY_LABEL.get(str(f.get("severity")), str(f.get("severity")))
        A(f"- **[{st_label}·{sev}] {f.get('title', '')}**（{f.get('id', '')}）")
        A(f"  - {f.get('detail', '')}")
        if f.get("numbers"):
            A(f"  - 关键数字：`{_json(f.get('numbers'))}`")
    A("")
    A("**局限说明**")
    A("")
    A("1. 论文表5.7/5.9 的判断矩阵无法复现表5.8/5.11 的权重，校准模式下系统并列展示论文基准与实时复算值。")
    A("2. 表5.17 的 Q 无法由表5.16 的 S/R 与式(4.11) 推出，Q 的对照仅作展示，不反向修正论文数值。")
    A("3. ISA 满意度存在表5.20 与图5.10 两套口径；实测情感口径由评论情感均值换算至李克特 1-5，"
      "与问卷口径不可直接等同。")
    A("4. 情感与评论统计随全局筛选变化，导出前请确认筛选口径与数据版本。")
    A("")

    # 九、附录
    A("## 九、附录")
    A("")
    A("**A. 场景参数（完整）**")
    A("")
    A("| 参数 | 取值 |")
    A("| --- | --- |")
    for k in ("mode", "lam", "v", "ideal_strategy", "weight_mode", "prob_completion",
              "direction", "use_paper_weights"):
        A(f"| {k} | {sc.get(k, '—')} |")
    A("")
    A("**B. 产物与路径**")
    A("")
    A(f"- 本报告：`{out_path}`")
    A(f"- 输出目录：`{REPORTS_DIR}`")
    A(f"- 报告标题：{context.get('title', DEFAULT_TITLE)}")
    A("")
    A("---")
    A("")
    A(f"*{context.get('title', DEFAULT_TITLE)} · 生成于 {ts}*")
    A("")

    out_path.write_text("\n".join(L), encoding="utf-8")
    return out_path


def _advice_body(record: dict) -> str:
    from core.algorithm.isa import quadrant_advice
    try:
        return quadrant_advice(record).get("body", "")
    except Exception:                                        # noqa: BLE001
        return ""


# =============================================================== PDF
@lru_cache(maxsize=1)
def _register_chinese_font() -> str:
    """注册可用中文字体：系统 TTF/TTC → CID STSong-Light 兜底。返回字体名。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/Supplemental/Songti.ttc",
        "/System/Library/Fonts/Supplemental/Kaiti.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/System/Library/Fonts/Supplemental/NotoSansCJK-Regular.otf",
    ]
    for i, path in enumerate(candidates):
        if not Path(path).exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont(f"DshCN{i}", path, subfontIndex=0))
            return f"DshCN{i}"
        except Exception:                                    # noqa: BLE001
            continue
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def pdf_probe() -> tuple[bool, str]:
    """探测 PDF 生成能力：(是否可用, 说明)。不写文件。"""
    try:
        import reportlab  # noqa: F401
    except ImportError as e:                                 # pragma: no cover
        return False, f"reportlab 未安装：{e}"
    try:
        font = _register_chinese_font()
        return True, f"reportlab {getattr(reportlab, '__version__', '')} · 中文字体 {font}"
    except Exception as e:                                   # noqa: BLE001  # pragma: no cover
        return False, f"中文字体注册失败：{type(e).__name__}: {e}"


def build_pdf(context: dict, out_path: Path) -> Path:
    """reportlab 中文 PDF：标题页 + 执行摘要 + 决策表 + ISA 表。

    reportlab 或中文字体不可用时抛 `ReportDependencyError`，由调用方降级为 Markdown。
    """
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                        Table, TableStyle)
    except ImportError as e:                                 # pragma: no cover
        raise ReportDependencyError(f"reportlab 不可用，无法生成 PDF：{e}") from e
    try:
        font = _register_chinese_font()
    except Exception as e:                                   # noqa: BLE001
        raise ReportDependencyError(f"中文字体注册失败，无法生成 PDF：{e}") from e

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    ts = _ts(context)
    sc = context.get("scenario", {}) or {}
    res = context.get("results", {}) or {}
    rows_r = res.get("rows", []) or []
    ranking = res.get("ranking", []) or []
    paper_rank = res.get("paper_ranking", []) or []
    audit = context.get("audit", {}) or {}
    summ = audit.get("summary", {}) or {}
    isa = context.get("isa", {}) or {}
    kpi = context.get("kpi", {}) or {}
    insight = context.get("insight", {}) or {}

    styles = {
        "title": ParagraphStyle("t", fontName=font, fontSize=24, leading=34,
                                alignment=TA_CENTER, textColor=colors.HexColor("#1F2937"),
                                spaceAfter=10),
        "subtitle": ParagraphStyle("s", fontName=font, fontSize=12, leading=20,
                                   alignment=TA_CENTER, textColor=colors.HexColor("#4B5563")),
        "h1": ParagraphStyle("h1", fontName=font, fontSize=15, leading=22,
                             textColor=colors.HexColor("#4F46E5"), spaceBefore=12, spaceAfter=6),
        "body": ParagraphStyle("b", fontName=font, fontSize=10.5, leading=17,
                               textColor=colors.HexColor("#111827"), spaceAfter=4),
        "muted": ParagraphStyle("m", fontName=font, fontSize=9, leading=14,
                                textColor=colors.HexColor("#6B7280")),
        "cell": ParagraphStyle("c", fontName=font, fontSize=9, leading=13,
                               textColor=colors.HexColor("#111827")),
        "cellhead": ParagraphStyle("ch", fontName=font, fontSize=9, leading=13,
                                   textColor=colors.white),
    }

    def P(text: str, style: str = "body") -> Paragraph:
        safe = (str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
        return Paragraph(safe, styles[style])

    def table(header: list[str], body: list[list[str]], widths=None) -> Table:
        data = [[P(h, "cellhead") for h in header]]
        data += [[P(c, "cell") for c in row] for row in body]
        t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F5F7FB")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return t

    story = []
    # -------- 标题页
    story.append(Spacer(1, 48 * mm))
    story.append(P(context.get("title", DEFAULT_TITLE), "title"))
    story.append(P("新能源汽车产品改进多属性决策 · 洞察与报告", "subtitle"))
    story.append(Spacer(1, 14 * mm))
    story.append(table(
        ["项目", "内容"],
        [["论文基准", f"{context.get('paper_title', '')}"],
         ["出处", context.get("thesis", "")],
         ["生成时间", ts],
         ["计算模式", f"{sc.get('mode_label', '—')}（λ={_f(sc.get('lam'), 2)}，v={_f(sc.get('v'), 2)}）"],
         ["数据范围", context.get("filters_desc", "全量数据（未筛选）")],
         ["数据版本", str(context.get("data_version") if context.get("data_version") is not None else "—")]],
        widths=[36 * mm, 120 * mm]))
    story.append(Spacer(1, 10 * mm))
    story.append(P("口径说明：本报告同时展示论文基准与平台实时复算值；"
                   "所有论文内部口径冲突在「审计与局限」中显式列出，系统不做静默修正。", "muted"))
    story.append(PageBreak())

    # -------- 执行摘要
    story.append(P("一、执行摘要", "h1"))
    if insight.get("headline"):
        story.append(P(str(insight["headline"]), "body"))
    kpi_rows = []
    if kpi:
        kpi_rows = [
            ["有效评论数", f"{int(kpi.get('n_kept', 0)):,}",
             f"保留率 {kpi.get('keep_rate', '—')}%"],
            ["筛选后情感均值", _f(kpi.get("mean_sentiment")),
             f"论文基准 {_f(kpi.get('paper_mean_sentiment'))}"],
            ["正/负评论", f"{int(kpi.get('pos', 0)):,} / {int(kpi.get('neg', 0)):,}",
             f"正负比 {kpi.get('pos_ratio', '—')}%"],
        ]
    kpi_rows.append(["复算 Top1", ranking[0] if ranking else "—",
                     f"排序：{' > '.join(ranking) if ranking else '—'}"])
    kpi_rows.append(["论文 Top1", paper_rank[0] if paper_rank else "—",
                     f"排序：{' > '.join(paper_rank) if paper_rank else '—'}"])
    if summ:
        kpi_rows.append(["一致性审计",
                         f"通过 {summ.get('pass', 0)} / 待确认 {summ.get('warn', 0)} / 偏差 {summ.get('fail', 0)}",
                         f"共 {summ.get('total', 0)} 项，高严重度 {summ.get('high', 0)} 项"])
    isa_imp = [r["属性"] for r in isa.get("records", []) if r.get("象限") == "改进区"]
    kpi_rows.append([f"ISA 改进区（{context.get('isa_source_label', '论文表5.20')}）",
                     "、".join(isa_imp) if isa_imp else "无", "改进优先级最高"])
    story.append(table(["指标", "数值", "备注"], kpi_rows, widths=[46 * mm, 52 * mm, 58 * mm]))
    story.append(Spacer(1, 5 * mm))
    story.append(P(f"排序一致性：{'完全一致' if ranking == paper_rank else '部分不一致'}"
                   f"（Kendall τ = {_kendall(ranking, paper_rank):.4f}）；"
                   f"λ 扫描稳定性：{'排序稳定' if (context.get('sensitivity', {}) or {}).get('lambda', {}).get('stable') else '存在排序翻转'}。",
                   "muted"))

    # -------- 决策表
    story.append(P("二、决策结果与排序（复算 vs 论文表5.16/5.17）", "h1"))
    if rows_r:
        body = [[r.get("属性", ""), str(r.get("复算排名", "—")), str(r.get("论文排名", "—")),
                 r.get("排序一致", "—"), _f(r.get("P(x)_复算")), _f(r.get("P(x)_论文")),
                 _f(r.get("ΔP(x)")), _f(r.get("Q_i")), _f(r.get("S_i")), _f(r.get("R_i"))]
                for r in rows_r]
        story.append(table(["属性", "复算名次", "论文名次", "一致", "P(x)复算", "P(x)论文",
                            "ΔP(x)", "Q", "S", "R"], body,
                           widths=[17 * mm, 16 * mm, 16 * mm, 12 * mm, 19 * mm, 19 * mm,
                                   17 * mm, 15 * mm, 15 * mm, 15 * mm]))
    else:
        story.append(P("（暂无决策结果）", "muted"))

    # -------- ISA 表
    story.append(P(f"三、ISA 改进优先级（口径：{context.get('isa_source_label', '论文表5.20')}）", "h1"))
    recs = isa.get("records", []) or []
    if recs:
        body = [[r.get("属性", ""), _f(r.get("重要性")), _f(r.get("满意度")),
                 r.get("象限", ""), _f(r.get("Δ重要性")), _f(r.get("Δ满意度")),
                 _advice_body(r)] for r in recs]
        story.append(table(["属性", "重要性", "满意度", "象限", "Δ重要性", "Δ满意度", "改进建议"], body,
                           widths=[16 * mm, 16 * mm, 16 * mm, 22 * mm, 18 * mm, 18 * mm, 50 * mm]))
    else:
        story.append(P("（暂无 ISA 数据）", "muted"))
    conflict = context.get("isa_conflict", []) or []
    if conflict:
        story.append(Spacer(1, 4 * mm))
        story.append(P("口径冲突提示：论文表5.20 与图5.10/正文对同一属性给出不同满意度"
                       "（详见审计项 ISA_SATISFACTION_CONFLICT）。", "muted"))

    # -------- 审计摘要
    story.append(P("四、审计与局限（摘要）", "h1"))
    story.append(P("以下问题来自论文数据本身，系统不做静默修正；完整明细见 Markdown / Excel 版本。", "muted"))
    body = [[STATUS_LABEL.get(str(f.get("status")), str(f.get("status"))),
             SEVERITY_LABEL.get(str(f.get("severity")), str(f.get("severity"))),
             str(f.get("title", ""))]
            for f in (audit.get("findings", []) or [])]
    if body:
        story.append(table(["状态", "严重度", "问题"], body,
                           widths=[24 * mm, 18 * mm, 114 * mm]))
    else:
        story.append(P("（暂无审计发现）", "muted"))
    story.append(Spacer(1, 6 * mm))
    story.append(P(f"生成于 {ts} · 输出目录 {REPORTS_DIR}", "muted"))

    doc = SimpleDocTemplate(str(out_path), pagesize=A4,
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title=context.get("title", DEFAULT_TITLE),
                            author="智评车行 · 决策看板")
    doc.build(story)
    return out_path


# =============================================================== 便捷入口
def new_output_path(kind: str) -> Path:
    """outputs/reports 下的时间戳文件路径（自动建目录、避免同秒覆盖）。"""
    if kind not in KIND_EXT:
        raise ValueError(f"未知报告类型：{kind}（可选 {list(KIND_EXT)}）")
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    base = REPORTS_DIR / f"洞察与报告_{stamp}_{kind}"
    path = base.with_suffix("." + KIND_EXT[kind])
    n = 1
    while path.exists():
        path = base.with_name(f"{base.name}_{n}").with_suffix("." + KIND_EXT[kind])
        n += 1
    return path


def build_payload(context: dict, kind: str) -> tuple[str, bytes]:
    """生成一种报告并返回 (文件名, bytes)，供下载按钮直接使用。"""
    builder = {"excel": build_excel, "markdown": build_markdown, "pdf": build_pdf}[kind]
    path = builder(context, new_output_path(kind))
    return path.name, path.read_bytes()
